"""Excel export: sorted by tier then score, frozen header, autofilter, colour-coded tiers."""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# (header, db field, width)
COLUMNS = [
    ("Tier", "tier", 6),
    ("Business name", "name", 32),
    ("Category", "category", 18),
    ("Group", "group_name", 16),
    ("Area", "area", 18),
    ("Address", "address", 38),
    ("Phone", "phone", 18),
    ("Email", "email", 30),
    ("Website", "website", 32),
    ("Score", "score", 7),
    ("Failed checks", "failed_checks", 50),
    ("Google rating", "rating", 8),
    ("Review count", "review_count", 8),
    ("Maps link", "maps_url", 14),
    ("Run filters", "run_filters", 40),
    ("Date found", "found_at", 12),
    ("Source", "_source", 22),
    ("Status", "outreach_status", 14),
    ("Notes", "notes", 40),
]

TIER_FILL = {
    "A": PatternFill("solid", fgColor="F8D7DA"),   # no real website: hottest
    "B": PatternFill("solid", fgColor="FDE5CC"),
    "C": PatternFill("solid", fgColor="FFF3C4"),
}
TIER_CELL = {
    "A": PatternFill("solid", fgColor="E57373"),
    "B": PatternFill("solid", fgColor="F0A35E"),
    "C": PatternFill("solid", fgColor="E8C547"),
}
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
STATUS_OPTIONS = ["New", "Contacted", "Replied", "Meeting", "Won", "Lost", "Do not contact"]
TIER_ORDER = {"A": 0, "B": 1, "C": 2}


def sort_leads(rows: list[dict]) -> list[dict]:
    def key(r):
        score = r.get("score")
        return (TIER_ORDER.get(r.get("tier"), 9), -1 if score is None or score != score else score, (r.get("name") or "").lower())

    return sorted(rows, key=key)


def _source_text(r: dict) -> str:
    src = {"google": "Google Places", "osm": "OpenStreetMap"}.get(r.get("source") or "", r.get("source") or "")
    if r.get("email"):
        via = {"osm": "OSM", "website": "website"}.get(r.get("email_source") or "", "")
        return f"{src}; email from {via}" if via else src
    return src


def export_leads(rows: list[dict], path: Path | str | BinaryIO) -> Path | BinaryIO:
    """Write to a file path, or to a file-like object (for the dashboard's download button)."""
    if isinstance(path, (str, Path)):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    for col, (header, _field, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 30

    for r_idx, r in enumerate(sort_leads(rows), start=2):
        tier = r.get("tier")
        for c_idx, (_header, field, _w) in enumerate(COLUMNS, start=1):
            value = _source_text(r) if field == "_source" else r.get(field)
            if value != value:  # NaN from pandas
                value = None
            if field == "found_at" and value:
                value = str(value)[:10]
            cell = ws.cell(row=r_idx, column=c_idx, value=value if value not in ("", None) else None)
            if tier in TIER_FILL:
                cell.fill = TIER_FILL[tier]
            if field in ("website", "maps_url") and value:
                cell.hyperlink = str(value)
                cell.font = Font(color="1A56DB", underline="single")
                if field == "maps_url":
                    cell.value = "Open map"
            if field == "failed_checks":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        tier_cell = ws.cell(row=r_idx, column=1)
        tier_cell.font = Font(bold=True)
        tier_cell.alignment = Alignment(horizontal="center")
        if tier in TIER_CELL:
            tier_cell.fill = TIER_CELL[tier]

    last_row = max(ws.max_row, 2)
    ws.freeze_panes = "C2"  # header row + business name stay visible
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"

    status_col = get_column_letter([c[1] for c in COLUMNS].index("outreach_status") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPTIONS) + '"', allow_blank=True)
    dv.add(f"{status_col}2:{status_col}{max(last_row, 500)}")
    ws.add_data_validation(dv)

    legend = wb.create_sheet("Key")
    legend.append(["Tier", "Meaning"])
    legend.append(["A", "No website, or no real website (social page only, unreachable, parked)"])
    legend.append(["B", "Website scores below 50"])
    legend.append(["C", "Website scores 50–70"])
    for i, t in enumerate("ABC", start=2):
        legend.cell(row=i, column=1).fill = TIER_CELL[t]
    legend.column_dimensions["B"].width = 70
    legend.cell(row=1, column=1).font = legend.cell(row=1, column=2).font = Font(bold=True)

    wb.save(path)
    return path


def default_export_path(export_dir: Path, suffix: str = "") -> Path:
    stamp = dt.datetime.now().strftime("%Y-%m-%d_%H%M")
    return Path(export_dir) / f"rome_leads_{stamp}{('_' + suffix) if suffix else ''}.xlsx"
